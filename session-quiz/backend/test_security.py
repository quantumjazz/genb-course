"""HTTP regression checks against disposable synthetic SQLite records only.

Run: python3 -m unittest -v test_security.py
No running server, production database or administrator secret is used.
"""

import io
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

import db
import server


class QuizSecurityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="quiz-security-")
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.db_path = root / "quiz.db"
        self.banks = root / "banks"
        self.banks.mkdir()
        # A real pre-update student schema exercises the additive migration.
        conn = db.connect(self.db_path)
        conn.execute("""CREATE TABLE quiz_student (
            session_id TEXT NOT NULL REFERENCES quiz_session(id) ON DELETE CASCADE,
            student_token TEXT NOT NULL, student_number TEXT NOT NULL,
            joined_at INTEGER NOT NULL, ended_at INTEGER, end_reason TEXT,
            PRIMARY KEY(session_id, student_token))""")
        db.init_schema(conn)
        with conn:
            conn.execute("INSERT INTO bank(id,name,uploaded_at) VALUES ('bank','Synthetic',1)")
            for i in range(10):
                conn.execute("""INSERT INTO bank_item
                    (id,bank_id,lecture_tag,stem,options_json,correct_index)
                    VALUES (?, 'bank', 'test', ?, ?, 0)""",
                    (f"item-{i}", f"Synthetic question {i}", '["Correct","Other"]'))
        conn.close()
        self.http = server.QuizServer(
            ("127.0.0.1", 0), server.Handler, str(self.db_path),
            "synthetic-admin-secret", None, self.banks, "", None,
        )
        thread = threading.Thread(target=self.http.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(self.close_server)
        self.base = f"http://127.0.0.1:{self.http.server_port}"
        self.admin = {"X-Admin-Key": "synthetic-admin-secret"}
        status, self.session, _ = self.request("POST", "/quiz/admin/session/create", {
            "bank_id": "bank", "lecture_tag": "test", "item_count": 5,
            "duration_minutes": 5, "security_mode": "strict",
        }, self.admin)
        self.assertEqual(status, 201, self.session)
        self.sid = self.session["session_id"]
        status, _, _ = self.request("POST", "/quiz/admin/session/start",
                                    {"session_id": self.sid}, self.admin)
        self.assertEqual(status, 200)

    def close_server(self):
        self.http.shutdown()
        self.http.server_close()
        for conn in self.http._thread_conns.values():
            conn.close()

    def request(self, method, path, body=None, headers=None, raw=False):
        headers = dict(headers or {})
        data = None
        if body is not None:
            data = json.dumps(body).encode()
            headers["Content-Type"] = "application/json"
        req = urllib.request.Request(self.base + path, data=data,
                                     headers=headers, method=method)
        try:
            response = urllib.request.urlopen(req, timeout=5)
        except urllib.error.HTTPError as exc:
            response = exc
        with response:
            content = response.read()
            result = content if raw else json.loads(content)
            return response.status, result, dict(response.headers)

    def issue(self, *numbers, replace=False):
        status, result, _ = self.request("POST", "/quiz/admin/session/access-codes", {
            "session_id": self.sid, "student_numbers": list(numbers),
            "replace": replace,
        }, self.admin)
        self.assertIn(status, (200, 201), result)
        return result

    def login(self, number, code):
        return self.request("POST", "/quiz/join", {
            "code": self.session["join_code"], "student_number": number,
            "access_code": code,
        })

    def student(self, number):
        issued = self.issue(number)
        code = issued["credentials"][0]["access_code"]
        status, result, _ = self.login(number, code)
        self.assertEqual(status, 200, result)
        return code, result["student_token"], {"X-Student-Token": result["student_token"]}

    def test_admin_access_fails_closed(self):
        for path in ("/quiz/admin/session/list", "/quiz/admin/bank/list"):
            self.assertEqual(self.request("GET", path)[0], 401)
        self.assertEqual(self.request("POST", "/quiz/admin/session/access-codes", {
            "session_id": self.sid, "student_numbers": ["TEST-001"],
        })[0], 401)
        self.http.admin_key = ""
        self.assertEqual(self.request("GET", "/quiz/admin/session/list", headers=self.admin)[0], 401)

    def test_private_code_required_and_secrets_not_stored_raw(self):
        code, token, headers = self.student("TEST-001")
        self.assertGreaterEqual(len(token), 40)
        for number, attempted in (("TEST-001", ""), ("TEST-001", "WRONG"), ("TEST-002", code)):
            self.assertEqual(self.login(number, attempted)[0], 401)
        self.assertEqual(self.request("GET", f"/quiz/next?student_token={token}")[0], 401)
        status, question, response_headers = self.request("GET", "/quiz/next", headers=headers)
        self.assertEqual(status, 200, question)
        self.assertIn("attempt_id", question)
        self.assertEqual(response_headers.get("Cache-Control"), "no-store")
        self.assertEqual(response_headers.get("Referrer-Policy"), "no-referrer")
        conn = db.connect(self.db_path)
        with conn:
            stored = conn.execute("SELECT access_code_hash FROM quiz_credential").fetchone()[0]
            bearer_hash = conn.execute("SELECT auth_token_hash FROM quiz_student").fetchone()[0]
        conn.close()
        self.assertNotEqual(stored, code)
        self.assertNotEqual(bearer_hash, token)

    def test_rejoin_revokes_old_bearer_and_preserves_question(self):
        code, old_token, headers = self.student("TEST-001")
        _, old_question, _ = self.request("GET", "/quiz/next", headers=headers)
        status, joined, _ = self.login("TEST-001", code)
        self.assertEqual(status, 200, joined)
        self.assertNotEqual(old_token, joined["student_token"])
        self.assertEqual(self.request("GET", "/quiz/next", headers=headers)[0], 401)
        new_headers = {"X-Student-Token": joined["student_token"]}
        _, question, _ = self.request("GET", "/quiz/next", headers=new_headers)
        self.assertEqual(question["attempt_id"], old_question["attempt_id"])

    def test_attempt_ownership_required_for_all_mutations(self):
        _, _, first_headers = self.student("TEST-001")
        _, _, other_headers = self.student("TEST-002")
        _, question, _ = self.request("GET", "/quiz/next", headers=first_headers)
        payloads = {
            "/quiz/answer": {"chosen_visible_index": 0},
            "/quiz/blur": {},
            "/quiz/incident": {"event_type": "window_blur"},
        }
        for path, extra in payloads.items():
            body = {"attempt_id": question["attempt_id"], **extra}
            self.assertEqual(self.request("POST", path, body)[0], 401, path)
            self.assertEqual(self.request("POST", path, body, other_headers)[0], 403, path)
        conn = db.connect(self.db_path)
        attempt = conn.execute("SELECT * FROM quiz_attempt WHERE id=?", (question["attempt_id"],)).fetchone()
        self.assertIsNone(attempt["submitted_at"])
        self.assertEqual(attempt["swapped"], 0)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM quiz_incident").fetchone()[0], 0)
        conn.close()

    def test_beacons_authenticated_and_device_metadata_discarded(self):
        _, token, headers = self.student("TEST-001")
        _, question, _ = self.request("GET", "/quiz/next", headers=headers)
        status, result, _ = self.request("POST", "/quiz/incident", {
            "student_token": token, "attempt_id": question["attempt_id"],
            "event_type": "window_blur", "client_ts": 9223372036854775807,
            "metadata": {
                "hidden": True, "fullscreen": False, "userAgent": "must-not-be-stored",
                "screenWidth": 12345, "arbitrary": {"data": "must-not-be-stored"},
            },
        })
        self.assertEqual(status, 201, result)
        conn = db.connect(self.db_path)
        metadata = json.loads(conn.execute("SELECT metadata_json FROM quiz_incident").fetchone()[0])
        self.assertIsNone(conn.execute("SELECT client_ts FROM quiz_incident").fetchone()[0])
        with conn:
            conn.execute("UPDATE quiz_incident SET client_ts=9223372036854775807")
        conn.close()
        self.assertEqual(metadata, {"hidden": True, "fullscreen": False})
        status, content, _ = self.request("GET", f"/quiz/admin/session/export?session_id={self.sid}",
                                         headers=self.admin, raw=True)
        self.assertEqual(status, 200)
        self.assertIn("incidents", load_workbook(io.BytesIO(content)).sheetnames)
        status, result, _ = self.request("POST", "/quiz/blur", {
            "student_token": token, "attempt_id": question["attempt_id"],
        })
        self.assertEqual(status, 200, result)
        self.assertTrue(result["swapped"])

    def test_code_reset_preserves_grade_and_revokes_access(self):
        code, _, headers = self.student("TEST-001")
        _, question, _ = self.request("GET", "/quiz/next", headers=headers)
        self.assertEqual(self.request("POST", "/quiz/answer", {
            "attempt_id": question["attempt_id"], "chosen_visible_index": 0,
        }, headers)[0], 200)
        duplicate = self.issue("TEST-001")
        self.assertEqual(duplicate["credentials"], [])
        self.assertIn("TEST-001", duplicate["skipped"])
        replacement = self.issue("TEST-001", replace=True)["credentials"][0]["access_code"]
        self.assertNotEqual(replacement, code)
        self.assertEqual(self.login("TEST-001", code)[0], 401)
        self.assertEqual(self.request("GET", "/quiz/next", headers=headers)[0], 401)
        self.assertEqual(self.login("TEST-001", replacement)[0], 200)
        status, content, _ = self.request("GET", f"/quiz/admin/session/export?session_id={self.sid}",
                                         headers=self.admin, raw=True)
        self.assertEqual(status, 200)
        workbook = load_workbook(io.BytesIO(content))
        self.assertEqual(workbook.sheetnames, ["summary", "detail", "incidents"])
        rows = list(workbook["summary"].values)
        self.assertEqual(rows[1][rows[0].index("student_number")], "TEST-001")
        self.assertEqual(rows[1][rows[0].index("items_answered")], 1)

    def test_legacy_pseudonym_rejected_but_records_survive_migration(self):
        conn = db.connect(self.db_path)
        session = server.session_by_id(conn, self.sid)
        legacy_token = server.derive_student_token(session["pseudonym_key"], "LEGACY-001")
        with conn:
            conn.execute("""INSERT INTO quiz_student
                (session_id,student_token,student_number,joined_at)
                VALUES (?,?,?,1)""", (self.sid, legacy_token, "LEGACY-001"))
            conn.execute("""INSERT INTO quiz_attempt
                (id,session_id,student_token,bank_item_id,ord,option_order_json,
                 chosen_index,correct,served_at,submitted_at,swapped)
                VALUES ('legacy-attempt',?,?,'item-0',1,'[0,1]',0,1,1,2,0)""",
                (self.sid, legacy_token))
        db.init_schema(conn)
        conn.close()
        self.assertEqual(self.request("GET", "/quiz/next",
                                     headers={"X-Student-Token": legacy_token})[0], 401)
        code = self.issue("LEGACY-001")["credentials"][0]["access_code"]
        status, joined, _ = self.login("LEGACY-001", code)
        self.assertEqual(status, 200, joined)
        conn = db.connect(self.db_path)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM quiz_student").fetchone()[0], 1)
        row = conn.execute("SELECT correct,submitted_at FROM quiz_attempt WHERE id='legacy-attempt'").fetchone()
        self.assertEqual(tuple(row), (1, 2))
        conn.close()

    def test_delete_cascades_credentials_and_student_records(self):
        self.student("TEST-001")
        status, result, _ = self.request("DELETE", f"/quiz/admin/session/{self.sid}", headers=self.admin)
        self.assertEqual(status, 200, result)
        conn = db.connect(self.db_path)
        for table in ("quiz_credential", "quiz_student", "quiz_attempt", "quiz_incident"):
            self.assertEqual(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0], 0, table)
        conn.close()


class StartupSecurityTests(unittest.TestCase):
    def test_startup_refuses_missing_admin_key_before_opening_database(self):
        with tempfile.TemporaryDirectory(prefix="quiz-startup-") as directory:
            env = dict(os.environ)
            env.pop("QUIZ_ADMIN_KEY", None)
            target = Path(directory) / "must-not-exist.db"
            env["QUIZ_DB_PATH"] = str(target)
            result = subprocess.run([sys.executable, str(Path(server.__file__)), "--port", "0"],
                                    env=env, capture_output=True, text=True, timeout=5)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("QUIZ_ADMIN_KEY", result.stdout + result.stderr)
            self.assertFalse(target.exists())


if __name__ == "__main__":
    unittest.main()
